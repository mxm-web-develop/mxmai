#!/usr/bin/env python3
"""
Platform IR → PPTX compiler (python-pptx).

Inspired by hugohe3/ppt-master (MIT) intermediate-language idea, but compiles
structured slide IR directly to native shapes.

Usage:
  python3 build_pptx_from_ir.py --input deck.json --output out.pptx
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
except ImportError as e:  # pragma: no cover
    print(f"ERROR: python-pptx required: {e}", file=sys.stderr)
    sys.exit(2)


def parse_hex(color: str | None, fallback: str) -> RGBColor:
    raw = (color or fallback).strip().lstrip("#")
    if len(raw) == 3:
        raw = "".join(c * 2 for c in raw)
    if len(raw) != 6:
        raw = fallback.lstrip("#")
    try:
        return RGBColor(int(raw[0:2], 16), int(raw[2:4], 16), int(raw[4:6], 16))
    except ValueError:
        return RGBColor(0x0F, 0x17, 0x2A)


def slide_size(ratio: str | None) -> tuple[int, int]:
    if (ratio or "16:9") == "4:3":
        return Inches(10), Inches(7.5)
    return Inches(13.333), Inches(7.5)


def write_paragraph(p, text: str, *, size_pt: float, bold: bool, color: RGBColor, font_name: str, align=None) -> None:
    if align is not None:
        p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name


def fill_bg(slide, prs, color: RGBColor) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    sp_tree = slide.shapes._spTree
    sp = shape._element
    sp_tree.remove(sp)
    sp_tree.insert(2, sp)


def render_slide(prs, slide_ir: dict[str, Any], vs: dict[str, Any]) -> None:
    palette = vs.get("palette") or {}
    typo = vs.get("typography") or {}
    bg = parse_hex(palette.get("background"), "#0F172A")
    fg = parse_hex(palette.get("foreground"), "#F8FAFC")
    accent = parse_hex(
        (slide_ir.get("page_style") or {}).get("accent") or palette.get("accent"),
        "#38BDF8",
    )
    muted = parse_hex(palette.get("muted"), "#94A3B8")
    title_font = str(typo.get("title_font") or "Microsoft YaHei")
    body_font = str(typo.get("body_font") or "Microsoft YaHei")
    title_size = float(typo.get("title_size_pt") or 36)
    body_size = float(typo.get("body_size_pt") or 18)

    blank = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
    slide = prs.slides.add_slide(blank)
    fill_bg(slide, prs, bg)

    layout = str(slide_ir.get("layout_hint") or "bullets")
    title = str(slide_ir.get("title") or "").strip()
    subtitle = str(slide_ir.get("subtitle") or "").strip()
    body = str(slide_ir.get("body") or "").strip()
    bullets = [str(b).strip() for b in (slide_ir.get("bullets") or []) if str(b).strip()]
    notes = str(slide_ir.get("notes") or "").strip()

    w = prs.slide_width
    h = prs.slide_height
    margin = Inches(0.7)

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), h)
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()

    if layout in ("title_center", "closing", "quote_center", "section_divider"):
        if title:
            box = slide.shapes.add_textbox(margin, Inches(2.4), w - 2 * margin, Inches(1.6))
            tf = box.text_frame
            tf.word_wrap = True
            write_paragraph(
                tf.paragraphs[0],
                title,
                size_pt=title_size + (8 if layout == "title_center" else 0),
                bold=True,
                color=fg,
                font_name=title_font,
                align=PP_ALIGN.CENTER,
            )
        if subtitle or body:
            text = subtitle or body
            box = slide.shapes.add_textbox(margin, Inches(4.2), w - 2 * margin, Inches(1.2))
            tf = box.text_frame
            tf.word_wrap = True
            write_paragraph(
                tf.paragraphs[0],
                text,
                size_pt=body_size,
                bold=False,
                color=muted,
                font_name=body_font,
                align=PP_ALIGN.CENTER,
            )
    else:
        if title:
            box = slide.shapes.add_textbox(margin, Inches(0.45), w - 2 * margin, Inches(1.0))
            tf = box.text_frame
            tf.word_wrap = True
            write_paragraph(
                tf.paragraphs[0], title, size_pt=title_size, bold=True, color=fg, font_name=title_font
            )
        if subtitle:
            box = slide.shapes.add_textbox(margin, Inches(1.35), w - 2 * margin, Inches(0.5))
            tf = box.text_frame
            write_paragraph(
                tf.paragraphs[0],
                subtitle,
                size_pt=body_size - 2,
                bold=False,
                color=muted,
                font_name=body_font,
            )

        top = Inches(2.0) if subtitle else Inches(1.7)
        content_h = h - top - Inches(0.6)

        if layout == "two_column" and len(bullets) >= 2:
            mid = max(1, len(bullets) // 2)
            for col, chunk in enumerate((bullets[:mid], bullets[mid:])):
                left = margin if col == 0 else w // 2 + Inches(0.2)
                width = w // 2 - margin - Inches(0.3)
                box = slide.shapes.add_textbox(left, top, width, content_h)
                tf = box.text_frame
                tf.word_wrap = True
                for i, item in enumerate(chunk):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    write_paragraph(
                        p, f"• {item}", size_pt=body_size, bold=False, color=fg, font_name=body_font
                    )
        elif layout == "three_cards" and bullets:
            cards = bullets[:3]
            gap = Inches(0.3)
            card_w = (w - 2 * margin - 2 * gap) // 3
            for i, item in enumerate(cards):
                left = margin + i * (card_w + gap)
                shape = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_w, content_h - Inches(0.4)
                )
                shape.fill.solid()
                shape.fill.fore_color.rgb = parse_hex(None, "#1E293B")
                shape.line.color.rgb = accent
                box = slide.shapes.add_textbox(
                    left + Inches(0.2),
                    top + Inches(0.3),
                    card_w - Inches(0.4),
                    content_h - Inches(0.8),
                )
                tf = box.text_frame
                tf.word_wrap = True
                write_paragraph(
                    tf.paragraphs[0], item, size_pt=body_size, bold=False, color=fg, font_name=body_font
                )
        elif layout == "big_number":
            num = bullets[0] if bullets else (body[:24] if body else title)
            box = slide.shapes.add_textbox(margin, Inches(2.5), w - 2 * margin, Inches(1.8))
            tf = box.text_frame
            write_paragraph(
                tf.paragraphs[0],
                num,
                size_pt=54,
                bold=True,
                color=accent,
                font_name=title_font,
                align=PP_ALIGN.CENTER,
            )
            rest = bullets[1:] if bullets else []
            text = " · ".join(rest) if rest else body
            if text:
                box2 = slide.shapes.add_textbox(margin, Inches(4.6), w - 2 * margin, Inches(1.5))
                tf2 = box2.text_frame
                tf2.word_wrap = True
                write_paragraph(
                    tf2.paragraphs[0],
                    text,
                    size_pt=body_size,
                    bold=False,
                    color=muted,
                    font_name=body_font,
                    align=PP_ALIGN.CENTER,
                )
        else:
            box = slide.shapes.add_textbox(margin, top, w - 2 * margin, content_h)
            tf = box.text_frame
            tf.word_wrap = True
            items = bullets if bullets else ([body] if body else ["（本页待补充内容）"])
            for i, item in enumerate(items):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                prefix = f"{i + 1}. " if layout == "agenda_list" else "• "
                write_paragraph(
                    p, f"{prefix}{item}", size_pt=body_size, bold=False, color=fg, font_name=body_font
                )

    if notes:
        try:
            slide.notes_slide.notes_text_frame.text = notes
        except Exception:
            pass


def build_presentation(deck: dict[str, Any]) -> Presentation:
    vs = deck.get("visual_system") or {}
    ratio = vs.get("ratio") or "16:9"
    prs = Presentation()
    prs.slide_width, prs.slide_height = slide_size(str(ratio))

    slides = deck.get("slides") or []
    if not slides:
        raise ValueError("slides[] empty")

    slides_sorted = sorted(slides, key=lambda s: int(s.get("order") or 0))
    for s in slides_sorted:
        render_slide(prs, s if isinstance(s, dict) else {}, vs)
    return prs


def build_pptx_bytes(deck: dict[str, Any]) -> bytes:
    from io import BytesIO

    prs = build_presentation(deck)
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build(deck: dict[str, Any], output: Path) -> None:
    prs = build_presentation(deck)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    slides = deck.get("slides") or []
    print(f"OK wrote {output} ({len(slides)} slides)")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build PPTX from deck IR JSON")
    parser.add_argument("--input", "-i", required=True)
    parser.add_argument("--output", "-o", required=True)
    args = parser.parse_args(argv)
    data = json.loads(Path(args.input).read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        print("ERROR: root must be object", file=sys.stderr)
        return 1
    try:
        build(data, Path(args.output))
    except ValueError as e:
        print(f"ERROR: {e}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
